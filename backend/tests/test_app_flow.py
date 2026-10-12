"""
End to end through the Flask app: seven answers in, a schedule (or an honest reason) out.
The Gemini / Claude SDKs talk real HTTP to tests/fake_llm.py.
"""

from __future__ import annotations

import io
import time

import pytest

from conftest import ANSWERS, CLAUDE_HEADERS, GEMINI_HEADERS, GEMINI_KEY, answer_all, wait_for_schedule

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def assert_good_week(data):
    assert data["status"] == "ready", data
    assert [d["day"] for d in data["schedule"]] == DAYS
    for day in data["schedule"]:
        assert len(day["blocks"]) >= 8
        for block in day["blocks"]:
            assert block["type"] in ("study", "break", "meal", "sleep", "class", "other")
            assert block["energy"] in ("high", "medium", "low", "")
            assert "–" in block["time"] and "AM" not in block["time"].upper().replace("MAIN", "")
    assert data["weekly_focus_hours"] > 0
    assert round(sum(s["hours_this_week"] for s in data["subject_breakdown"]), 1) == data["weekly_focus_hours"]
    assert data["reasoning_notes"]
    assert data["meta"]["provider"] in ("gemini", "claude")


# ---------------------------------------------------------------------------
# The headline scenario: the 7 questions, a Gemini key, and the schedule comes back
# ---------------------------------------------------------------------------
def test_seven_answers_with_gemini_key_returns_schedule(client, fake_llm):
    responses = answer_all(client, GEMINI_HEADERS)
    assert [r.status_code for r in responses] == [200] * 7
    first = responses[0].get_json()
    assert first["done"] is False and first["step"] == 1 and first["total"] == 7
    assert first["reply"].startswith("Got it, thanks!")
    last = responses[-1].get_json()
    assert last["done"] is True and last["step"] == 7
    assert "building your week" in last["reply"].lower()
    assert last["schedule_status"] in ("generating", "ready")

    data = wait_for_schedule(client)
    assert_good_week(data)
    assert data["meta"]["provider"] == "gemini"

    # The right models were used and the key never travelled in a body.
    chat_models = set(fake_llm.models_called("gemini", "chat"))
    assert chat_models == {"gemini-3.5-flash-lite"}                    # fast tier
    schedule_calls = fake_llm.calls("gemini", "schedule")
    assert [c["model"] for c in schedule_calls] == ["gemini-3.8-flash"]   # balanced tier, one call
    assert schedule_calls[0]["thinking_level"] in ("low", "LOW") and schedule_calls[0]["thinking_budget"] is None
    assert schedule_calls[0]["max_tokens"] == 12288


def test_effort_tiers_pick_models_and_token_caps(client, fake_llm):
    assert client.post("/api/effort", json={"level": "high"}).get_json() == {"status": "ok", "level": "high"}
    answer_all(client, GEMINI_HEADERS)
    assert_good_week(wait_for_schedule(client))
    call = fake_llm.calls("gemini", "schedule")[0]
    assert call["max_tokens"] == 16384
    assert call["thinking_level"] in ("medium", "MEDIUM")


def test_low_effort_returns_compact_schedule(client, fake_llm):
    client.post("/api/effort", json={"level": "low"})
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    assert all(b["note"] == "" for d in data["schedule"] for b in d["blocks"])
    assert fake_llm.calls("gemini", "schedule")[0]["max_tokens"] == 8192


def test_claude_key_path(client, fake_llm):
    answer_all(client, CLAUDE_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    assert data["meta"]["provider"] == "claude"
    assert fake_llm.models_called("claude", "schedule") == ["claude-sonnet-5-5"]
    assert all(not c["has_temperature"] for c in fake_llm.calls("claude"))
    assert all(c["first_role"] == "user" for c in fake_llm.calls("claude"))


# ---------------------------------------------------------------------------
# The failure that started all this: models that no longer exist, thinking tokens, quotas, keys
# ---------------------------------------------------------------------------
def test_retired_models_fall_through_to_one_that_works(client, fake_llm):
    fake_llm.scenario.unavailable_models = {"gemini-3.8-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite"}
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    tried = fake_llm.models_called("gemini", "schedule")
    assert tried[:3] == ["gemini-3.8-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite"] or "gemini-flash-latest" in tried
    assert data["meta"]["model"] not in {"gemini-3.8-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite"}


def test_every_configured_model_gone_uses_model_discovery(client, fake_llm):
    configured = {"gemini-3.5-flash-lite", "gemini-3.1-flash-lite", "gemini-3.8-flash", "gemini-flash-lite-latest",
                  "gemini-flash-latest", "gemini-2.5-flash", "gemini-3.6-flash"}
    fake_llm.scenario.unavailable_models = configured
    fake_llm.scenario.gemini_models = ["gemini-3.9-flash", "gemini-3.7-flash-lite", "embedding-001"]
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    assert data["meta"]["model"] == "gemini-3.9-flash"


def test_thinking_tokens_eating_the_budget_are_survived(client, fake_llm):
    fake_llm.scenario.thinking_tokens = 9000        # real Gemini counts these against max_output_tokens
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    caps = [c["max_tokens"] for c in fake_llm.calls("gemini", "schedule")]
    assert len(caps) == 2 and caps[1] > caps[0]      # first call came back cut off, second got more room


def test_small_chat_cap_never_returns_an_empty_reply(client, fake_llm):
    fake_llm.scenario.thinking_tokens = 5000          # would swallow the old max_tokens=200 entirely
    r = client.post("/api/chat", json={"message": ANSWERS[0]}, headers=GEMINI_HEADERS).get_json()
    assert r["reply"].startswith("Got it, thanks!")
    assert fake_llm.calls("gemini", "chat")[0]["max_tokens"] == 1024


def test_model_that_rejects_thinking_level_is_retried_without_it(client, fake_llm):
    fake_llm.scenario.reject_thinking_level = {"gemini-3.8-flash"}
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    calls = fake_llm.calls("gemini", "schedule")
    assert calls[0]["model"] == "gemini-3.8-flash" and calls[0]["thinking_level"] is not None
    assert calls[1]["model"] == "gemini-3.8-flash" and calls[1]["thinking_level"] is None and calls[1]["thinking_budget"] is None


def test_rate_limit_with_retry_hint_recovers(client, fake_llm):
    fake_llm.scenario.queue = [{"provider": "gemini", "status": 429, "api_status": "RESOURCE_EXHAUSTED",
                                "message": "Quota exceeded for metric: requests per minute", "retry_after": 2}]
    r = client.post("/api/chat", json={"message": ANSWERS[0]}, headers=GEMINI_HEADERS).get_json()
    assert r["reply"].startswith("Got it")            # waited the 2 s (patched) and retried the same model
    assert "notice" not in r


def test_daily_quota_exhausted_is_reported_not_hidden(client, fake_llm):
    fake_llm.scenario.always_fail = {"provider": "gemini", "status": 429, "api_status": "RESOURCE_EXHAUSTED",
                                     "message": "You exceeded your current quota. Quota exceeded for metric: "
                                                "generate_content_free_tier_requests, limit: 0"}
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert data["status"] == "failed" and data["reason"] == "quota_exhausted"
    assert "no quota left" in data["error"]
    assert data["schedule"] == []


def test_bad_key_surfaces_at_question_one_and_in_the_final_status(client, fake_llm):
    fake_llm.scenario.valid_gemini_keys = {"AIzaSyTHE_REAL_KEY_THAT_IS_LONG_ENOUGH_12345"}
    first = client.post("/api/chat", json={"message": ANSWERS[0]}, headers=GEMINI_HEADERS).get_json()
    assert first["notice"]["reason"] == "bad_key"
    assert first["reply"] == "Which country is that in?"      # the plain question: the student is never stuck
    rest = answer_all(client, GEMINI_HEADERS, ANSWERS[1:])
    assert rest[-1].get_json()["done"] is True
    data = wait_for_schedule(client)
    assert data["status"] == "failed" and data["reason"] == "bad_key"
    assert "didn't accept this API key" in data["error"]
    assert GEMINI_KEY not in str(data)


def test_fixing_the_key_then_retrying_builds_the_week_without_redoing_the_chat(client, fake_llm):
    good = "AIzaSyTHE_REAL_KEY_THAT_IS_LONG_ENOUGH_12345"
    fake_llm.scenario.valid_gemini_keys = {good}
    answer_all(client, GEMINI_HEADERS)
    assert wait_for_schedule(client)["reason"] == "bad_key"
    started = client.post("/api/schedule/retry", headers={"X-Gemini-Api-Key": good}).get_json()
    assert started["status"] == "generating"
    assert_good_week(wait_for_schedule(client))


def test_no_key_anywhere_fails_fast_with_no_key(client, fake_llm):
    responses = answer_all(client, {})
    assert responses[0].get_json()["notice"]["reason"] == "no_key"
    assert responses[0].get_json()["reply"] == "Which country is that in?"
    assert responses[-1].get_json()["done"] is True
    data = wait_for_schedule(client)
    assert data["status"] == "failed" and data["reason"] == "no_key"
    assert fake_llm.calls() == []                      # nothing was ever sent to a provider


def test_server_side_key_is_used_when_the_student_has_none(client, fake_llm, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "AIzaSyOPERATOR_OWNED_KEY_0123456789abcd")
    monkeypatch.setenv("DEFAULT_LLM_PROVIDER", "gemini")
    cfg = client.get("/api/config").get_json()
    assert cfg["server_llm"] == {"available": True, "provider": "gemini", "providers": ["gemini"]}
    answer_all(client, {})
    data = wait_for_schedule(client)
    assert_good_week(data)
    assert "AIzaSyOPERATOR" not in str(cfg) and "AIzaSyOPERATOR" not in str(data)


def test_legacy_provider_value_does_not_crash(client, monkeypatch):
    monkeypatch.setenv("DEFAULT_LLM_PROVIDER", "legacy-provider")
    assert client.get("/api/config").get_json()["server_llm"]["provider"] is None


def test_provider_outage_is_reported_as_provider_error(client, fake_llm):
    fake_llm.scenario.always_fail = {"provider": "gemini", "status": 503, "api_status": "UNAVAILABLE",
                                     "message": "The model is overloaded. Please try again later."}
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert data["status"] == "failed" and data["reason"] == "provider_error"
    assert "hiccup" in data["error"]


def test_safety_block_is_reported(client, fake_llm):
    fake_llm.scenario.safety_block = True
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert data["status"] == "failed" and data["reason"] == "blocked"


def test_region_block_is_reported(client, fake_llm):
    fake_llm.scenario.region_blocked = True
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert data["status"] == "failed" and data["reason"] == "region_blocked"
    assert "Claude key" in data["error"]


# ---------------------------------------------------------------------------
# Sloppy but recoverable model output
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("style", ["ampm", "fenced", "dict_days", "messy_types"])
def test_sloppy_schedule_replies_are_repaired(client, fake_llm, style):
    fake_llm.scenario.schedule_style = style
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    assert len(fake_llm.calls("gemini", "schedule")) == 1          # repaired by code, no second paid call
    assert any(b["time"] == "9:00–11:00" for b in data["schedule"][0]["blocks"])
    if style == "messy_types":
        assert {b["type"] for d in data["schedule"] for b in d["blocks"]} <= {"study", "break", "meal", "sleep", "class", "other"}


def test_a_week_missing_two_days_gets_one_corrective_retry(client, fake_llm):
    fake_llm.scenario.broken_first_schedule = "five_days"
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    assert len(fake_llm.calls("gemini", "schedule")) == 2
    assert data["warnings"] == []


def test_a_week_that_stays_incomplete_is_shown_with_a_warning(client, fake_llm):
    fake_llm.scenario.schedule_style = "five_days"
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert data["status"] == "ready" and len(data["schedule"]) == 5
    assert any("Saturday and Sunday" in w for w in data["warnings"])


def test_empty_replies_fail_with_bad_response_not_a_crash(client, fake_llm):
    fake_llm.scenario.schedule_style = "empty_text"
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert data["status"] == "failed" and data["reason"] == "bad_response"


# ---------------------------------------------------------------------------
# Uploads feed the structuring step
# ---------------------------------------------------------------------------
def _docx_bytes(lines):
    import docx
    buffer = io.BytesIO()
    document = docx.Document()
    for line in lines:
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Operating Systems", "4 credits"
    table.cell(1, 0).text, table.cell(1, 1).text = "Linear Algebra", "3 credits"
    document.save(buffer)
    return buffer.getvalue()


def test_uploaded_syllabus_shapes_the_week(client, fake_llm):
    body = _docx_bytes(["Semester plan", "Data Structures - 3 lectures a week", "Midterm Exam: 2026-11-20"])
    up = client.post("/api/upload", data={"file": (io.BytesIO(body), "syllabus.docx")},
                     content_type="multipart/form-data").get_json()
    assert up == {"status": "received", "filename": "syllabus.docx"}
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert_good_week(data)
    subjects = {s["subject"] for s in data["subject_breakdown"]}
    assert {"Data Structures", "Operating Systems", "Linear Algebra"} <= subjects
    assert data["meta"]["documents"] == 1 and data["meta"]["subjects_found"] == 3
    assert fake_llm.calls("gemini", "extract")                    # the structuring step really ran
    assert fake_llm.calls("gemini", "extract")[0]["max_tokens"] == 4096


def test_unreadable_documents_still_build_a_week_and_say_so(client, fake_llm):
    body = _docx_bytes(["Data Structures"])
    client.post("/api/upload", data={"file": (io.BytesIO(body), "syllabus.docx")}, content_type="multipart/form-data")
    original = fake_llm._content

    def broken_extract(task, system, last_user):
        return "I cannot help with that." if task == "extract" else original(task, system, last_user)
    fake_llm._content = broken_extract
    try:
        answer_all(client, GEMINI_HEADERS)
        data = wait_for_schedule(client)
    finally:
        fake_llm._content = original
    assert data["status"] == "ready"
    assert any("couldn't read your uploaded files" in w for w in data["warnings"])


def test_bad_key_while_structuring_stops_early_without_a_second_failing_call(client, fake_llm):
    body = _docx_bytes(["Data Structures"])
    client.post("/api/upload", data={"file": (io.BytesIO(body), "s.docx")}, content_type="multipart/form-data")
    fake_llm.scenario.valid_gemini_keys = {"AIzaSyTHE_REAL_KEY_THAT_IS_LONG_ENOUGH_12345"}
    answer_all(client, GEMINI_HEADERS)
    data = wait_for_schedule(client)
    assert data["reason"] == "bad_key"
    assert fake_llm.calls("gemini", "schedule") == []


# ---------------------------------------------------------------------------
# The job is idempotent, pollable and resettable
# ---------------------------------------------------------------------------
def test_status_is_generating_while_the_build_runs(client, fake_llm):
    fake_llm.scenario.latency_s = 0.4
    answer_all(client, GEMINI_HEADERS)
    seen = set()
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        data = client.get("/api/schedule").get_json()
        seen.add(data["status"])
        if data["status"] == "generating":
            assert data["progress"]["stage"] in ("queued", "structuring", "scheduling")
            assert data["schedule"] == []
        if data["status"] in ("ready", "failed"):
            break
        time.sleep(0.05)
    assert "generating" in seen and data["status"] == "ready"


def test_sending_another_message_after_done_does_not_start_a_second_build(client, fake_llm):
    answer_all(client, GEMINI_HEADERS)
    wait_for_schedule(client)
    again = client.post("/api/chat", json={"message": "one more thing"}, headers=GEMINI_HEADERS).get_json()
    assert again["done"] is True and again["schedule_status"] == "ready"
    assert len(fake_llm.calls("gemini", "schedule")) == 1


def test_retry_never_replaces_a_good_schedule_or_runs_twice(client, fake_llm):
    answer_all(client, GEMINI_HEADERS)
    first = wait_for_schedule(client)
    again = client.post("/api/schedule/retry", headers=GEMINI_HEADERS).get_json()
    assert again["status"] == "ready" and again["schedule"] == first["schedule"]
    assert len(fake_llm.calls("gemini", "schedule")) == 1


def test_retry_before_the_chat_is_finished_does_nothing(client):
    client.post("/api/chat", json={"message": ANSWERS[0]}, headers=GEMINI_HEADERS)
    assert client.post("/api/schedule/retry", headers=GEMINI_HEADERS).get_json()["status"] == "pending"


def test_chat_state_restores_progress_after_a_refresh(client):
    client.post("/api/chat", json={"message": ANSWERS[0]}, headers=GEMINI_HEADERS)
    client.post("/api/chat", json={"message": ANSWERS[1]}, headers=GEMINI_HEADERS)
    state = client.get("/api/chat/state").get_json()
    assert state["step"] == 2 and state["total"] == 7 and state["done"] is False
    assert [m["role"] for m in state["messages"]] == ["user", "assistant", "user", "assistant"]
    assert state["greeting"].endswith("What college or university do you attend?")


def test_reset_starts_over(client, fake_llm):
    answer_all(client, GEMINI_HEADERS)
    wait_for_schedule(client)
    assert client.post("/api/reset").get_json() == {"status": "ok"}
    state = client.get("/api/chat/state").get_json()
    assert state["step"] == 0 and state["messages"] == [] and state["schedule_status"] == "pending"
    assert client.get("/api/schedule").get_json()["status"] == "pending"


def test_two_browsers_have_separate_sessions(app_module, fake_llm):
    a, b = app_module.app.test_client(), app_module.app.test_client()
    answer_all(a, GEMINI_HEADERS)
    assert wait_for_schedule(a)["status"] == "ready"
    assert b.get("/api/schedule").get_json()["status"] == "pending"
    assert b.get("/api/chat/state").get_json()["step"] == 0


def test_empty_message_keeps_the_original_400_shape(client):
    r = client.post("/api/chat", json={"message": "   "}, headers=GEMINI_HEADERS)
    assert r.status_code == 400
    body = r.get_json()
    assert body["reply"] == "Could you type something first?" and body["done"] is False


# ---------------------------------------------------------------------------
# /api/key/test
# ---------------------------------------------------------------------------
def test_key_test_ok(client, fake_llm):
    body = client.post("/api/key/test", headers=GEMINI_HEADERS).get_json()
    assert body["ok"] is True and body["provider"] == "gemini" and body["model"] == "gemini-3.5-flash-lite"
    assert body["source"] == "user" and "key" not in str(body).lower().replace("source", "")


def test_key_test_bad_key(client, fake_llm):
    fake_llm.scenario.valid_gemini_keys = {"AIzaSyTHE_REAL_KEY_THAT_IS_LONG_ENOUGH_12345"}
    body = client.post("/api/key/test", headers=GEMINI_HEADERS).get_json()
    assert body["ok"] is False and body["reason"] == "bad_key" and "didn't accept" in body["message"]


def test_key_test_without_a_key(client):
    body = client.post("/api/key/test").get_json()
    assert body["ok"] is False and body["reason"] == "no_key"


def test_key_with_impossible_shape_is_rejected_before_any_network_call(client, fake_llm):
    body = client.post("/api/key/test", headers={"X-Gemini-Api-Key": "has spaces inside it and is long enough!!"}).get_json()
    assert body["reason"] == "bad_key"
    assert fake_llm.calls() == []


# ---------------------------------------------------------------------------
# Hardening
# ---------------------------------------------------------------------------
def test_rate_limit_returns_429_with_retry_after(client, app_module, monkeypatch):
    monkeypatch.setitem(app_module.LIMITS, "chat", (3, 100))
    codes = [client.post("/api/chat", json={"message": "x"}, headers={}).status_code for _ in range(5)]
    assert codes[:3] == [200, 200, 200] and codes[3:] == [429, 429]
    blocked = client.post("/api/chat", json={"message": "x"})
    assert blocked.get_json()["reason"] == "too_many_requests" and int(blocked.headers["Retry-After"]) >= 1


def test_oversized_json_body_is_refused(client):
    r = client.post("/api/chat", json={"message": "x" * 70_000})
    assert r.status_code == 413 and r.get_json()["reason"] == "too_large"


def test_overlong_answer_is_clipped_not_rejected(client):
    r = client.post("/api/chat", json={"message": "a" * 5000}, headers={})
    assert r.status_code == 200
    state = client.get("/api/chat/state").get_json()
    assert len(state["messages"][0]["content"]) == 2000


def test_unknown_api_path_and_wrong_method_answer_in_json(client):
    assert client.get("/api/nope").get_json()["reason"] == "not_found"
    r = client.get("/api/chat")
    assert r.status_code == 405 and r.get_json()["reason"] == "method_not_allowed"


def test_security_headers_and_cookie_flags(client):
    r = client.get("/api/chat/state")
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert "frame-ancestors 'none'" in r.headers["Content-Security-Policy"]
    assert "script-src 'self'" in r.headers["Content-Security-Policy"]
    assert r.headers["Cache-Control"] == "no-store"
    cookie = r.headers["Set-Cookie"]
    assert "HttpOnly" in cookie and "SameSite=Lax" in cookie


def test_api_keys_never_appear_in_any_response_or_log(client, fake_llm, caplog):
    import logging
    caplog.set_level(logging.DEBUG)
    answer_all(client, GEMINI_HEADERS)
    wait_for_schedule(client)
    client.post("/api/key/test", headers=GEMINI_HEADERS)
    state = client.get("/api/chat/state").get_json()
    assert GEMINI_KEY not in str(state) and GEMINI_KEY not in caplog.text


def test_static_pages_are_served_and_path_traversal_is_not(client):
    assert client.get("/").status_code == 200
    assert client.get("/chat.html").status_code == 200
    assert client.get("/results.html").status_code == 200
    assert client.get("/../backend/app.py").status_code == 404
    assert client.get("/%2e%2e/backend/app.py").status_code == 404


def test_health_and_config(client):
    assert client.get("/api/health").get_json()["status"] == "ok"
    cfg = client.get("/api/config").get_json()
    assert cfg["total_questions"] == 7 and cfg["max_upload_mb"] == 20 and "pdf" in cfg["allowed_extensions"]
    assert [p["id"] for p in cfg["providers"]] == ["gemini", "claude"]
    assert "groq" not in str(cfg).lower() and "free" not in str(cfg).lower()
