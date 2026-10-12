"""
Unit tests for document_parser.py: byte sniffing, zip-bomb guard, format parsing,
the .xls-that-is-really-.xlsx fallback, and friendly error messages.
"""

from __future__ import annotations

import io
import os
import zipfile

import docx
import openpyxl
import pytest
from PIL import Image

import document_parser as dp


def write(tmp_path, name: str, data: bytes) -> str:
    path = os.path.join(tmp_path, name)
    with open(path, "wb") as f:
        f.write(data)
    return path


# ---------------------------------------------------------------------------
# sniff_kind
# ---------------------------------------------------------------------------
class TestSniffKind:
    def test_pdf_header(self, tmp_path):
        path = write(tmp_path, "a.pdf", b"%PDF-1.4\n%fake")
        assert dp.sniff_kind(path) == "pdf"

    def test_zip_header(self, tmp_path):
        path = write(tmp_path, "a.docx", b"PK\x03\x04" + b"\x00" * 20)
        assert dp.sniff_kind(path) == "zip"

    def test_ole_header_for_legacy_doc(self, tmp_path):
        path = write(tmp_path, "a.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 20)
        assert dp.sniff_kind(path) == "ole"

    def test_png_header(self, tmp_path):
        path = write(tmp_path, "a.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 20)
        assert dp.sniff_kind(path) == "png"

    def test_jpeg_header(self, tmp_path):
        path = write(tmp_path, "a.jpg", b"\xff\xd8\xff\xe0" + b"\x00" * 20)
        assert dp.sniff_kind(path) == "jpeg"

    def test_unrecognised_bytes_return_none(self, tmp_path):
        path = write(tmp_path, "a.pdf", b"this is just plain text, not a pdf at all")
        assert dp.sniff_kind(path) is None

    def test_renamed_executable_is_not_a_pdf(self, tmp_path):
        path = write(tmp_path, "totally-a.pdf", b"MZ\x90\x00\x03\x00\x00\x00" + b"\x00" * 100)
        assert dp.sniff_kind(path) != "pdf"


class TestParseDocumentRejectsMismatchedBytes:
    def test_exe_renamed_to_pdf_is_refused(self, tmp_path):
        path = write(tmp_path, "evil.pdf", b"MZ\x90\x00\x03\x00\x00\x00" + b"\x00" * 200)
        result = dp.parse_document(path, "evil.pdf")
        assert result.success is False
        assert "doesn't look like" in result.error

    def test_text_file_renamed_to_docx_is_refused(self, tmp_path):
        path = write(tmp_path, "fake.docx", b"just some plain text content here")
        result = dp.parse_document(path, "fake.docx")
        assert result.success is False

    def test_unsupported_extension_is_refused_before_any_read(self, tmp_path):
        path = write(tmp_path, "a.exe", b"MZ\x90\x00")
        result = dp.parse_document(path, "a.exe")
        assert result.success is False and "Unsupported" in result.error


# ---------------------------------------------------------------------------
# .xls that is really .xlsx (renamed / saved with the old extension)
# ---------------------------------------------------------------------------
class TestXlsXlsxFallback:
    def test_xls_with_pk_header_is_read_as_xlsx(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Subject", "Hours"])
        ws.append(["Data Structures", 3])
        buf = io.BytesIO()
        wb.save(buf)
        path = write(tmp_path, "renamed.xls", buf.getvalue())
        result = dp.parse_document(path, "renamed.xls")
        assert result.success is True
        assert "Data Structures" in result.text and "3" in result.text

    def test_genuine_xls_via_xlrd(self, tmp_path):
        xlwt = pytest.importorskip("xlwt")
        wb = xlwt.Workbook()
        ws = wb.add_sheet("Sheet1")
        ws.write(0, 0, "Subject")
        ws.write(1, 0, "Linear Algebra")
        buf = io.BytesIO()
        wb.save(buf)
        path = write(tmp_path, "legacy.xls", buf.getvalue())
        result = dp.parse_document(path, "legacy.xls")
        assert result.success is True and "Linear Algebra" in result.text


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------
class TestDocx:
    def test_paragraphs_and_table_rows_in_order(self, tmp_path):
        document = docx.Document()
        document.add_paragraph("Semester Plan")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Data Structures", "3 credits"
        document.add_paragraph("Exam: 2026-11-20")
        path = os.path.join(tmp_path, "s.docx")
        document.save(path)
        result = dp.parse_document(path, "s.docx")
        assert result.success is True
        assert "Semester Plan" in result.text
        assert "Data Structures" in result.text and "3 credits" in result.text
        assert "Exam: 2026-11-20" in result.text
        assert result.text.index("Semester Plan") < result.text.index("Data Structures") < result.text.index("Exam")

    def test_doc_extension_that_is_really_docx_is_parsed(self, tmp_path):
        document = docx.Document()
        document.add_paragraph("Hello")
        path = os.path.join(tmp_path, "renamed.doc")
        document.save(path)
        result = dp.parse_document(path, "renamed.doc")
        assert result.success is True and "Hello" in result.text

    def test_genuine_legacy_doc_is_refused_with_a_clear_message(self, tmp_path):
        path = write(tmp_path, "old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100)
        result = dp.parse_document(path, "old.doc")
        assert result.success is False and ".docx" in result.error


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------
class TestXlsx:
    def test_multiple_sheets_and_dates(self, tmp_path):
        import datetime
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Courses"
        ws1.append(["Operating Systems", 4])
        ws2 = wb.create_sheet("Exams")
        ws2.append(["Midterm", datetime.date(2026, 11, 20)])
        path = os.path.join(tmp_path, "plan.xlsx")
        wb.save(path)
        result = dp.parse_document(path, "plan.xlsx")
        assert result.success is True
        assert "[Sheet: Courses]" in result.text and "[Sheet: Exams]" in result.text
        assert "Operating Systems" in result.text
        assert "2026-11-20" in result.text

    def test_hidden_sheet_is_skipped(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.append(["visible data"])
        hidden = wb.create_sheet("Secret")
        hidden.append(["hidden data"])
        hidden.sheet_state = "hidden"
        path = os.path.join(tmp_path, "h.xlsx")
        wb.save(path)
        result = dp.parse_document(path, "h.xlsx")
        assert "visible data" in result.text and "hidden data" not in result.text

    def test_empty_workbook_reports_no_readable_text(self, tmp_path):
        wb = openpyxl.Workbook()
        path = os.path.join(tmp_path, "empty.xlsx")
        wb.save(path)
        result = dp.parse_document(path, "empty.xlsx")
        assert result.success is False and "No readable text" in result.error


# ---------------------------------------------------------------------------
# Zip-bomb / corrupted archive guard (applies to docx and xlsx, both zip containers)
# ---------------------------------------------------------------------------
class TestZipGuard:
    def test_oversized_unpacked_content_is_refused(self, tmp_path, monkeypatch):
        monkeypatch.setattr(dp, "MAX_UNPACKED_BYTES", 1000)
        path = os.path.join(tmp_path, "bomb.docx")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("big.bin", b"0" * 5000)
        result = dp.parse_document(path, "bomb.docx")
        assert result.success is False and "larger than it looks" in result.error

    def test_too_many_entries_is_refused(self, tmp_path, monkeypatch):
        monkeypatch.setattr(dp, "MAX_ZIP_ENTRIES", 5)
        path = os.path.join(tmp_path, "many.docx")
        with zipfile.ZipFile(path, "w") as zf:
            for i in range(10):
                zf.writestr(f"f{i}.txt", "x")
        result = dp.parse_document(path, "many.docx")
        assert result.success is False

    def test_corrupted_zip_gives_a_friendly_message(self, tmp_path):
        path = write(tmp_path, "broken.docx", b"PK\x03\x04" + b"not actually a valid zip stream" * 5)
        result = dp.parse_document(path, "broken.docx")
        assert result.success is False
        assert "damaged" in result.error.lower() or "doesn't look like" in result.error.lower()


# ---------------------------------------------------------------------------
# Images / OCR path (tesseract availability checked, not OCR accuracy)
# ---------------------------------------------------------------------------
class TestImageHandling:
    def test_tiny_blank_image_reports_no_text_or_is_handled_gracefully(self, tmp_path):
        img = Image.new("RGB", (200, 80), color="white")
        path = os.path.join(tmp_path, "blank.png")
        img.save(path)
        result = dp.parse_document(path, "blank.png")
        # Either tesseract is missing (friendly message) or it runs and finds nothing on a blank image.
        assert result.success is False
        assert result.error

    def test_image_with_mismatched_bytes_is_refused(self, tmp_path):
        path = write(tmp_path, "fake.png", b"this is not a real png file at all")
        result = dp.parse_document(path, "fake.png")
        assert result.success is False


# ---------------------------------------------------------------------------
# Extension helpers
# ---------------------------------------------------------------------------
class TestGetExtension:
    @pytest.mark.parametrize("name,ext", [("a.pdf", "pdf"), ("A.PDF", "pdf"), ("my.file.docx", "docx"),
                                          ("noext", ""), ("", ""), (".hidden", "hidden")])
    def test_cases(self, name, ext):
        assert dp.get_extension(name) == ext
