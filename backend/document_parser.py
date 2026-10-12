"""
document_parser.py
Script 2 of 4 in the pipeline.

Job: take an uploaded file (PDF, Word, Excel, or image) and extract its raw
text. This script does NOT call any LLM: it is plain deterministic code.
(Worth remembering for Q&A: this is the one pipeline stage that isn't "AI",
it's extraction libraries doing exactly what they're told.)

Images go through OCR via pytesseract, and so do scanned PDFs (a PDF with no
real text layer): the first few pages are rendered and read, within a hard
time budget, and the student is told the text was recognised from a scan.
OCR quality on messy handwriting or low-quality photos is NOT guaranteed; the
note returned with such a file says so, and this is a named limitation, not
something to oversell in a demo.

Requires the system package `tesseract-ocr` for OCR (not just the Python
wrapper); the Dockerfile installs it. Without it, text files still work and
scanned files fail with a plain-language message.

Uploads are untrusted input, so before any parser sees a file:
  * the first bytes must match the claimed type (a renamed .exe is not a .pdf);
  * Word/Excel files (zip containers) are checked for zip-bomb proportions;
  * parsers are capped (pages, sheets, rows, characters) so one file cannot
    exhaust memory or time;
  * errors shown to the student are friendly sentences; the technical
    exception text only goes to the server log.
"""

from __future__ import annotations

import logging
import os
import time
import warnings
import zipfile
from datetime import date, datetime
from datetime import time as dtime

import docx
import openpyxl
import pdfplumber
import pytesseract
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {"pdf", "doc", "docx", "xls", "xlsx", "png", "jpg", "jpeg", "webp"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}

MAX_TEXT_CHARS = 20000          # stored per document; the structurer sends even less
MAX_PDF_PAGES = 60
MAX_SHEETS = 12
MAX_ROWS_PER_SHEET = 2000
MIN_PDF_TEXT_CHARS = 40         # fewer characters than this across all pages = "probably a scan"
OCR_PDF_MAX_PAGES = max(0, int(os.getenv("OCR_PDF_MAX_PAGES", "5")))  # 0 turns the scanned-PDF fallback off
OCR_BUDGET_S = 40               # total time one file may spend in OCR
OCR_PAGE_TIMEOUT_S = 20
OCR_DPI = 200
MAX_IMAGE_SIDE = 2600           # larger photos are shrunk before OCR (faster, and plenty for printed text)
MAX_UNPACKED_BYTES = 120 * 1024 * 1024
MAX_ZIP_ENTRIES = 3000

Image.MAX_IMAGE_PIXELS = 60_000_000  # Pillow's decompression-bomb guard; its warning is turned into an error below


class ParseError(Exception):
    """An expected failure; the message is written for the student."""


class ParseResult:
    """Simple result wrapper so callers always get a consistent shape back."""

    def __init__(self, success, text="", error=None, source_filename="", note=None, truncated=False):
        self.success = success
        self.text = text
        self.error = error
        self.source_filename = source_filename
        self.note = note              # something the student should know about a successful read
        self.truncated = truncated

    def to_dict(self):
        return {
            "success": self.success,
            "text": self.text,
            "error": self.error,
            "source_filename": self.source_filename,
            "note": self.note,
            "truncated": self.truncated,
        }


def get_extension(filename):
    name = str(filename or "").strip()
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


# ---------------------------------------------------------------------------
# Looking at the bytes, not the file name
# ---------------------------------------------------------------------------
_ALLOWED_KINDS = {
    "pdf": {"pdf"},
    "docx": {"zip"},
    "xlsx": {"zip"},
    "doc": {"ole", "zip"},   # a .doc is old OLE; "zip" = a .docx saved under the old extension
    "xls": {"ole", "zip"},   # same story for .xls vs .xlsx
    **{e: {"png", "jpeg", "webp"} for e in IMAGE_EXTENSIONS},
}


def sniff_kind(filepath) -> str | None:
    """Identify a file by its first bytes: pdf | zip | ole | png | jpeg | webp | None."""
    with open(filepath, "rb") as f:
        head = f.read(1024)
    if b"%PDF-" in head:
        return "pdf"
    if head.startswith((b"PK\x03\x04", b"PK\x05\x06")):
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "ole"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    return None


def _check_zip(filepath) -> None:
    """Word and Excel files are zip containers: refuse ones that unpack to absurd sizes."""
    try:
        with zipfile.ZipFile(filepath) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNPACKED_BYTES:
                raise ParseError("This file is far larger than it looks once unpacked, so it was not opened.")
    except zipfile.BadZipFile as exc:
        raise ParseError("This file looks damaged: it isn't a valid Word or Excel file.") from exc


# ---------------------------------------------------------------------------
# OCR helpers
# ---------------------------------------------------------------------------
_tesseract_ok = False


def _require_tesseract(what: str) -> None:
    """Fail with a plain sentence when the tesseract program is missing on this machine."""
    global _tesseract_ok
    if _tesseract_ok:
        return
    try:
        pytesseract.get_tesseract_version()
    except pytesseract.TesseractNotFoundError as exc:
        raise ParseError(f"{what} needs text recognition, which isn't installed on this server. "
                         "Try a text-based PDF, or a Word or Excel file.") from exc
    _tesseract_ok = True


def _ocr(image: Image.Image, timeout_s: float) -> str:
    try:
        return pytesseract.image_to_string(image, timeout=timeout_s)
    except RuntimeError as exc:  # pytesseract raises RuntimeError("Tesseract process timeout")
        raise ParseError("Reading this took too long. Try a smaller or clearer file.") from exc


def _prepare_for_ocr(image: Image.Image) -> Image.Image:
    """Upright, flattened onto white (no transparency), greyscale and at a sane size."""
    image = ImageOps.exif_transpose(image)
    if image.mode not in ("L", "RGB"):
        rgba = image.convert("RGBA")
        white = Image.new("RGBA", rgba.size, (255, 255, 255, 255))
        image = Image.alpha_composite(white, rgba)
    image = image.convert("L")
    image.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    return image


# ---------------------------------------------------------------------------
# Format parsers
# ---------------------------------------------------------------------------
def parse_pdf(filepath, *, allow_ocr: bool = True):
    """
    Extract text from a PDF. A real text layer is read directly with pdfplumber. If the
    pages hold (almost) no text it is probably a scan, so the first OCR_PDF_MAX_PAGES
    pages are rendered and read with OCR within OCR_BUDGET_S seconds.

    Returns (text, note); note is None or a sentence for the student.
    """
    note = None
    with pdfplumber.open(filepath) as pdf:
        total_pages = len(pdf.pages)
        pages = pdf.pages[:MAX_PDF_PAGES]
        parts = []
        for page in pages:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
        text = "\n".join(parts).strip()
        if total_pages > MAX_PDF_PAGES:
            note = f"Only the first {MAX_PDF_PAGES} pages were read."

        if len(text) >= MIN_PDF_TEXT_CHARS or not allow_ocr or OCR_PDF_MAX_PAGES == 0:
            return text, note

        _require_tesseract("This PDF looks scanned, so reading it")
        started = time.monotonic()
        ocr_parts, read = [], 0
        for page in pages[:OCR_PDF_MAX_PAGES]:
            if time.monotonic() - started > OCR_BUDGET_S:
                break
            image = _prepare_for_ocr(page.to_image(resolution=OCR_DPI).original)
            ocr_parts.append(_ocr(image, OCR_PAGE_TIMEOUT_S))
            read += 1
        ocr_text = "\n".join(p.strip() for p in ocr_parts if p.strip()).strip()
        if ocr_text:
            of = f" of {total_pages}" if total_pages > read else ""
            note = (f"This PDF was a scan, so we read {read} page{'s' if read != 1 else ''}{of} with text "
                    "recognition. Double-check the subjects and dates in your schedule.")
        logger.info("pdf OCR: %d/%d pages, %d chars, %.1fs", read, total_pages, len(ocr_text), time.monotonic() - started)
        return ocr_text, note


def _cell_texts(row, cap: int = 60) -> str:
    """One table row as 'a | b | c'; a merged cell is listed once, not once per column it spans."""
    seen, texts = set(), []
    for cell in row.cells[:cap]:
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        texts.append(cell.text.strip())
    line = " | ".join(texts)
    return line if line.strip(" |") else ""


def parse_docx(filepath):
    """Extract text from a Word (.docx) file: paragraphs and table rows, in document order."""
    document = docx.Document(filepath)
    text_parts = []
    rows_seen = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document).text.strip()
            if paragraph:
                text_parts.append(paragraph)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                rows_seen += 1
                if rows_seen > MAX_ROWS_PER_SHEET:
                    break
                line = _cell_texts(row)
                if line:
                    text_parts.append(line)
    return "\n".join(text_parts).strip()


def _fmt_cell(value) -> str:
    """Spreadsheet cell -> text that keeps dates readable (an exam date is the whole point)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == dtime(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dtime):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_xlsx(filepath):
    """
    Extract text from an Excel file: every visible sheet, every non-empty row.
    `filepath` may be a path or an open binary file.
    """
    workbook = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    text_parts = []
    try:
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            if getattr(sheet, "sheet_state", "visible") != "visible":
                continue
            sheet.reset_dimensions()  # some exporters write a wrong size, which would hide most of the sheet
            text_parts.append(f"[Sheet: {sheet.title}]")
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_number > MAX_ROWS_PER_SHEET:
                    break
                values = [t for t in (_fmt_cell(c) for c in row) if t]
                if values:
                    text_parts.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(text_parts).strip()


def parse_xls(filepath):
    """
    Extract text from a legacy Excel (.xls) file via xlrd (openpyxl only reads .xlsx).
    Same output shape as parse_xlsx(). Dates are converted from Excel serial numbers to
    ISO dates so exam dates stay readable for the syllabus-structuring step.

    Files named .xls are sometimes really .xlsx (renamed, or saved by tools that use the
    old extension). .xlsx is a zip container, so sniff the first bytes and hand those to
    parse_xlsx() instead of failing.
    """
    with open(filepath, "rb") as f:
        if f.read(2) == b"PK":
            f.seek(0)
            # Pass the open file, not the path: openpyxl rejects any *path*
            # ending in ".xls" by name, but accepts a file object.
            return parse_xlsx(f)

    import xlrd  # imported here so a missing xlrd only affects .xls uploads

    workbook = xlrd.open_workbook(filepath, on_demand=True)
    text_parts = []
    try:
        for sheet in workbook.sheets()[:MAX_SHEETS]:
            if getattr(sheet, "visibility", 0) != 0:
                continue
            text_parts.append(f"[Sheet: {sheet.name}]")
            for row_idx in range(min(sheet.nrows, MAX_ROWS_PER_SHEET)):
                row_values = []
                for cell in sheet.row(row_idx):
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode).isoformat(sep=" ")
                        except (xlrd.XLDateError, ValueError, OverflowError):
                            pass  # keep the raw number rather than dropping the cell
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                        value = int(value)  # 3.0 -> 3
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    text = str(value).strip()
                    if text:
                        row_values.append(text)
                if row_values:
                    text_parts.append(" | ".join(row_values))
    finally:
        workbook.release_resources()
    return "\n".join(text_parts).strip()


def parse_image(filepath):
    """OCR an image file using pytesseract."""
    _require_tesseract("Reading an image")
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        try:
            with Image.open(filepath) as image:
                prepared = _prepare_for_ocr(image)
        except (Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
            raise ParseError("This image is too large to read. Try a smaller photo.") from exc
    return _ocr(prepared, OCR_PAGE_TIMEOUT_S * 2).strip()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def _friendly_error(exc: Exception) -> str:
    low = str(exc).lower()
    name = type(exc).__name__
    if "password" in low or "encrypt" in low or name in ("PDFPasswordIncorrect", "PDFEncryptionError"):
        return "This file is password-protected. Remove the password and upload it again."
    if isinstance(exc, zipfile.BadZipFile) or name in ("PackageNotFoundError", "InvalidFileException",
                                                      "PDFSyntaxError", "XLRDError", "UnidentifiedImageError"):
        return "This file looks damaged, so it couldn't be read."
    return "Couldn't read this file. It may be damaged or password-protected."


def parse_document(filepath, original_filename=None):
    """
    Main entry point. Looks at the file extension, checks the bytes agree with it, and
    routes to the right parser. Returns a ParseResult: always check .success before using .text.
    """
    display_name = original_filename or os.path.basename(filepath)
    ext = get_extension(display_name)

    if ext not in SUPPORTED_EXTENSIONS:
        return ParseResult(success=False, error=f"Unsupported file type: .{ext}", source_filename=display_name)

    try:
        kind = sniff_kind(filepath)
        if kind not in _ALLOWED_KINDS[ext]:
            raise ParseError(f"This file doesn't look like a real .{ext} file. "
                             "It may be renamed, damaged or password-protected.")
        if kind == "zip":
            _check_zip(filepath)

        note = None
        if ext == "pdf":
            text, note = parse_pdf(filepath)
        elif ext == "docx":
            text = parse_docx(filepath)
        elif ext == "doc":
            if kind != "zip":
                # Legacy .doc (not .docx) is not supported by python-docx.
                raise ParseError("Old .doc format isn't supported — please upload as .docx.")
            text = parse_docx(filepath)  # a .docx saved under the old extension
        elif ext == "xls":
            text = parse_xls(filepath)
        elif ext == "xlsx":
            text = parse_xlsx(filepath)
        else:
            text = parse_image(filepath)

        text = text.strip()
        if not text:
            hint = " Try a clearer photo or a text-based file." if ext in IMAGE_EXTENSIONS or kind == "pdf" else ""
            raise ParseError("No readable text found in this file." + hint)

        truncated = len(text) > MAX_TEXT_CHARS
        if truncated:
            text = text[:MAX_TEXT_CHARS]
            note = (note + " " if note else "") + "This file is long, so only the first part is used."
        logger.info("Parsed %s (%s): %d chars, truncated=%s", display_name, ext, len(text), truncated)
        return ParseResult(success=True, text=text, source_filename=display_name, note=note, truncated=truncated)

    except ParseError as exc:
        logger.info("Could not parse %s: %s", display_name, exc)
        return ParseResult(success=False, error=str(exc), source_filename=display_name)
    except Exception as exc:  # noqa: BLE001 - parsers throw many things; the student gets a sentence, the log gets the detail
        logger.warning("Failed to parse %s: %s: %s", display_name, type(exc).__name__, str(exc)[:300])
        return ParseResult(success=False, error=_friendly_error(exc), source_filename=display_name)
